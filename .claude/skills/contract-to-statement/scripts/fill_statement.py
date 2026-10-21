"""extract_contract.py의 JSON으로 거래명세서 양식(xlsx)을 채워 저장한다.

사용법:
    python fill_statement.py <계약.json> [--month 2026-10] [--stage 1] [--set key=value ...]
                             [--out-dir C:\\ChaAI\\07_auto_doc\\거래명세서]

처리 흐름:
    1. --set 값으로 JSON의 빈 항목을 보완한다(예: supplier.phone=02-000-0000).
    2. 청구 월(--month, 기본값은 이번 달)의 말일을 발행일로 삼는다. 간이형 계약서의
       '매월 말일을 작성일자로' 규칙과 같다.
    3. 금액 행을 하나 만든다.
       - 월 금액 계약: 품명에 'YYYY년 M월분'을 붙이고, 단위 월, 수량 1, 단가 월 금액.
       - 분할 지급 계약: --stage로 고른 차수 금액. 차수를 고르지 않으면 오류를 낸다(어느 차수인지 추측하지 않음).
       - 총액 계약: 총 금액 1식.
       - 부가세 포함 금액이면 공급가액 = 금액 / 1.1을 원 단위로 반올림한다. 양식의 부가세 수식(K30*10%)이 나머지를 계산한다.
    4. 양식의 수식·병합·서식은 그대로 두고 값 칸만 채운다. 예시 값(H11=1, I11=50000)은 지운다.
    5. 같은 이름의 파일이 있으면 덮어쓰지 않고 _2, _3을 붙인다.

종료 코드 2: 필수 값이 비어 있다. 출력된 키를 사용자에게 받아 --set으로 다시 실행한다.
"""
import argparse
import calendar
import datetime as dt
import json
import re
import shutil
import sys
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

import openpyxl

SKILL_DIR = Path(__file__).resolve().parent.parent
TEMPLATE = SKILL_DIR / 'assets' / '거래명세서_양식.xlsx'
DEFAULT_OUT = Path(r'C:\ChaAI\07_auto_doc\거래명세서')

# 비어 있으면 거래명세서를 만들 수 없는 항목. 나머지(전화·주소·계좌·규격)는 사용자가 비워 두기로 할 수 있다.
REQUIRED = ('supplier.name', 'supplier.biz', 'buyer.name', 'amount', 'item.name')


def set_path(data, key, value):
    """'supplier.phone' 같은 점 경로에 값을 넣는다. 사용자가 '-'를 주면 빈칸 유지(의도적 생략)로 기록한다."""
    target = data
    *parents, last = key.split('.')
    for p in parents:
        target = target.setdefault(p, {})
    target[last] = '' if value == '-' else value


def get_path(data, key):
    """점 경로의 값을 읽는다. 없으면 None."""
    for p in key.split('.'):
        if not isinstance(data, dict):
            return None
        data = data.get(p)
    return data


def supply_value(value, vat_included):
    """부가세 포함 금액이면 공급가액(÷1.1, 원 단위 반올림)으로, 별도 금액이면 그대로 돌려준다."""
    if not vat_included:
        return value
    return int((Decimal(value) / Decimal('1.1')).quantize(Decimal('1'), rounding=ROUND_HALF_UP))


def build_row(data, month, stage):
    """품목 한 행(품명, 규격, 단위, 수량, 단가)과 비고 문구를 만든다."""
    item, amount, installments = data['item'], data['amount'], data.get('installments') or []
    label = f'{month.year}년 {month.month}월분'
    if installments:
        if stage is None:
            raise SystemExit('분할 지급 계약입니다. --stage 로 차수를 지정하세요: '
                             + ', '.join(f"{i['stage']}차 {i['value']:,}원" for i in installments))
        inst = next((i for i in installments if i['stage'] == stage), None)
        if inst is None:
            raise SystemExit(f'{stage}차분이 계약서에 없습니다.')
        name = f"{item['name']} {stage}차분({inst['percent']}%)"
        price = supply_value(inst['value'], inst['vat_included'])
        return name, item.get('spec'), item.get('unit') or '식', 1, price, f'{stage}차분'
    if amount['basis'] == '월':
        return (f"{item['name']} ({label})", item.get('spec'), item.get('unit') or '월', 1,
                supply_value(amount['value'], amount['vat_included']), label)
    return (item['name'], item.get('spec'), item.get('unit') or '식', 1,
            supply_value(amount['value'], amount['vat_included']), '')


def period_warning(data, issue):
    """발행일이 계약기간 밖이면 경고 문구를 돌려준다. 기간을 모르면 None."""
    p = data.get('period') or {}
    if not p.get('start') or not p.get('end'):
        return None
    if not (p['start'] <= issue.isoformat() <= p['end']):
        return f"발행일 {issue} 이(가) 계약기간 {p['start']}~{p['end']} 밖입니다."
    return None


def unique_path(folder, stem):
    """덮어쓰지 않도록 같은 이름이 있으면 _2, _3을 붙인 경로를 돌려준다."""
    path = folder / f'{stem}.xlsx'
    n = 2
    while path.exists():
        path = folder / f'{stem}_{n}.xlsx'
        n += 1
    return path


def safe(text):
    """파일 이름에 쓸 수 없는 문자를 없앤다."""
    return re.sub(r'[\\/:*?"<>|\s]+', '', text or '')


def fill(data, month, stage, out_dir):
    """양식을 채워 저장하고 (경로, 행 정보)를 돌려준다."""
    last_day = calendar.monthrange(month.year, month.month)[1]
    issue = dt.date(month.year, month.month, last_day)
    name, spec, unit, qty, price, note = build_row(data, month, stage)
    sup, buyer = data['supplier'], data['buyer']

    wb = openpyxl.load_workbook(TEMPLATE)
    ws = wb.active
    ws['B4'] = issue                                     # 셀 서식이 yyyy"년" m"월" d"일"
    ws['A6'] = f"{buyer['name']}  귀하"
    ws['H5'] = sup.get('biz') or ''
    ws['H6'] = sup.get('name') or ''
    ws['M6'] = f"{sup.get('rep') or ''} (인)".strip()
    ws['H7'] = sup.get('phone') or ''
    ws['H8'] = sup.get('address') or ''
    for r in range(11, 29):                              # 예시 값 제거
        for c in ('A', 'B', 'G', 'H', 'I'):
            ws[f'{c}{r}'] = None
    ws['A11'] = issue                                    # 셀 서식이 m/d
    ws['B11'] = f'{name} / {spec}' if spec else name
    ws['G11'] = unit
    ws['H11'] = qty
    ws['I11'] = price
    remarks = [f"계약번호 {data.get('contract_id') or '-'}"] + ([note] if note else [])
    ws['B30'] = ' · '.join(remarks)
    ws['A33'] = f"     계좌번호 : {data.get('account') or ''}"
    wb.calculation.fullCalcOnLoad = True                 # 엑셀이 열 때 수식을 다시 계산하게 한다

    out_dir.mkdir(parents=True, exist_ok=True)
    stem = f"거래명세서_{safe(data.get('contract_id')) or '계약번호없음'}_{safe(sup.get('name'))}_{month:%Y-%m}"
    if stage:
        stem += f'_{stage}차'
    path = unique_path(out_dir, stem)
    wb.save(path)
    vat = int((Decimal(price * qty) * Decimal('0.1')))
    return path, {'issue': str(issue), 'item': ws['B11'].value, 'unit': unit, 'qty': qty, 'price': price,
                  'supply': price * qty, 'vat_formula_result': vat, 'total': price * qty + vat,
                  'warning': period_warning(data, issue)}


def main():
    """인자를 읽고 필수 값을 확인한 뒤 양식을 채운다. 결과 요약은 UTF-8 JSON으로도 남긴다."""
    ap = argparse.ArgumentParser(description='계약 JSON → 거래명세서 xlsx')
    ap.add_argument('json')
    ap.add_argument('--month', help='청구 월 YYYY-MM (기본: 이번 달)')
    ap.add_argument('--stage', type=int, help='분할 지급 차수')
    ap.add_argument('--set', action='append', default=[], metavar='KEY=VALUE',
                    help="빈 값 보완. '-'는 빈칸으로 두기")
    ap.add_argument('--out-dir', default=str(DEFAULT_OUT))
    args = ap.parse_args()

    data = json.loads(Path(args.json).read_text(encoding='utf-8'))
    for pair in args.set:
        key, _, value = pair.partition('=')
        set_path(data, key.strip(), value.strip())
    month = dt.datetime.strptime(args.month, '%Y-%m').date() if args.month else dt.date.today().replace(day=1)

    lacking = [k for k in REQUIRED if not get_path(data, k)]
    if lacking:
        print('필수 값 없음: ' + ', '.join(lacking))
        return 2

    path, row = fill(data, month, args.stage, Path(args.out_dir))
    summary = {'saved': str(path), **row,
               'blank': [k for k in data.get('missing', []) if not get_path(data, k)]}
    Path(args.json).with_suffix('.result.json').write_text(
        json.dumps(summary, ensure_ascii=False, indent=2), encoding='utf-8')
    print(f"saved {path}")
    return 0


if __name__ == '__main__':
    # Windows 콘솔 기본 인코딩(cp949) 대신 UTF-8로 출력해 한글 메시지가 깨지지 않게 한다.
    sys.stdout.reconfigure(encoding='utf-8')
    sys.stderr.reconfigure(encoding='utf-8')
    sys.exit(main())
